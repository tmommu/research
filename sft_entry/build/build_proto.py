import re, sys
sys.path.insert(0, '.')
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_COLOR_INDEX, WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import content as C

OUT = sys.argv[1] if len(sys.argv) > 1 else "../SAGIP_SFT2026_Prototype_Design.docx"
d = docx.Document("templates/tplproto_prototype_template.docx")
body = d.element.body
for el in list(body):
    if el.tag != qn('w:sectPr'):
        body.remove(el)
BLUE = RGBColor(0x0B, 0x5D, 0x8F)
TOK = re.compile(r'(\*\*.+?\*\*|\[[^\]]+\])')

def add_text(p, text, size=None, bold=None, italic=None):
    for part in TOK.split(text):
        if not part: continue
        if part.startswith('**'):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith('['):
            r = p.add_run(part); r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            if bold: r.bold = True
        else:
            r = p.add_run(part)
            if bold: r.bold = True
        if size: r.font.size = Pt(size)
        if italic: r.italic = True

def title(t):
    p = d.add_paragraph(); r = p.add_run(t); r.bold = True; r.font.size = Pt(17); r.font.color.rgb = BLUE
def h(t):
    p = d.add_paragraph(); p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(3); p.paragraph_format.keep_with_next = True
    r = p.add_run(t); r.bold = True; r.font.size = Pt(13); r.font.color.rgb = BLUE
def para(t, **k):
    p = d.add_paragraph(); add_text(p, t, **k); return p
def bullet(label, t):
    p = d.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(1)
    r = p.add_run(label); r.bold = True
    add_text(p, t)
def fig(path, w, cap):
    d.add_picture(path, width=Inches(w)); d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(cap); r.italic = True; r.font.size = Pt(8.5)
def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement('w:shd')
    s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), fill); tcPr.append(s)
def table(rows, widths, bold_rows=()):
    t = d.add_table(rows=len(rows), cols=len(widths)); t.style = 'Table Grid'; t.autofit = False
    for ci, w in enumerate(widths): t.columns[ci].width = Inches(w)
    for ri, row in enumerate(rows):
        for ci, v in enumerate(row):
            c = t.cell(ri, ci); c.width = Inches(widths[ci])
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(1)
            add_text(p, v, size=9.5, bold=(ri == 0 or ri in bold_rows))
            if ri == 0: shade(c, 'DCEBF5')
            if ci == len(widths) - 1 and re.fullmatch(r'[+\d,]+( to [\d,]+)?( \(.*\))?', v or ''): p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    d.add_paragraph().paragraph_format.space_after = Pt(0)

title(C.TITLE + " (Prototype Design)")
para("Quick visual of what we will build and buy. Quezon City Science High School, Grade 10 team. Adviser: Skjoukje U. Jemoya. Site: Dario River, [Barangay Santo Cristo, Bago Bantay, to confirm]. All costs and sizes are planning estimates.")

h("1. What it is")
para("A solar-powered robot boat, about 1.7 m long, 0.85 m wide and 0.55 m high with a draft near 10 cm, navigates on its own along a preset route between two floating barriers in a creek. A bow camera runs an object detector that boxes floating plastic, the boat steers to it, "
     "and a scoop lifts it onto a cleated conveyor, one item per pocket. A second camera under a light-tight hood names each item with a small 2D convolutional neural network (CNN), "
     "and a servo gate sorts plastic into Bin A and plants, mixed or unsure items into Bin B. A weekly offline report goes to the barangay and a monthly Plastic Trend Scoreboard goes up "
     "at the Barangay Hall. A separate float-switch node on the bank sounds a flood alarm.")
fig("fig_boat_side.png", 6.6, "Side view of the SAGIP boat (schematic).")
fig("fig_patrol.png", 6.6, "Patrol, detect, scoop and dock (top view, schematic).")
fig("fig_3d_iso.png", 6.6, "3D model, isometric view with the deck see-through. Numbers match the parts table. Side and top views are in SAGIP_Boat_3D_Model.png, and the STL imports into Tinkercad.")

h("2. Parts and what each does")
table([("No.", "Part", "What it does"),
       ("1", "Bow camera (ESP32-CAM Type C)", "Runs the object detector and boxes floating plastic ahead so the boat can steer to it"),
       ("2", "Scoop", "Lifts debris onto the belt"),
       ("3", "Cleated conveyor belt", "Carries items up, one item per cleat pocket"),
       ("4", "Belt camera and CNN (ESP32-S3) under a hood", "Names each item in the same light, and a blob count flags mixed items"),
       ("5", "Servo gate and two bins", "Plastic to Bin A, plants, mixed or unsure items to Bin B"),
       ("6", "Electronics box", "ESP32 boards, drivers, converters, tilt sensor"),
       ("7", "30 W solar panel", "Flat on the deck, charges the battery"),
       ("8", "12 V 12 Ah sealed lead-acid battery", "Powers thrusters, belt and boards, kept low between the floats"),
       ("9", "Two thrusters", "Steer by differential thrust"),
       ("10", "Two side ultrasonic sensors and bump switches", "Keep the boat off the bank and sense contact at the bow"),
       ("11", "Twin 6-inch PVC floats and deck", "Keep it afloat and stable"),
       ("12", "Flood node on the bank", "Three float switches, buzzer and LED, protected 18650 cell in an IP67 box")],
      [0.5, 2.5, 3.4])

h("3. Why it is built this way")
para("Water hyacinth mats carried most of the floating plastic counted in a Saigon River study (Schreyers et al., 2021), so the upstream barrier angles mats to the bank, the detector separates plants and the propellers are shrouded. "
     "Floating plastic transport rose up to six times at peak river flow (van Emmerik et al., 2022), so the boat logs every item with time and the flood node logs the water level.")
para("INT8 quantization shrinks a model about fourfold (Gholami et al., 2021) and TensorFlow Lite for Microcontrollers runs it on an ESP32 (David et al., 2021). "
     "The creek banks and two barriers define a preset route, so the boat needs only side sensors and bump switches, with no GPS or compass. "
     "Twin floats with a wide beam, a low battery, foam in the bow and stern and a tilt cut-off keep it stable, and float and tilt tests come first.")
para("A skirted black hood and a fixed LED light every photo the same way, and a hydrophobic film, an angled window and a desiccant pack keep camera housings clear of droplets and fog. "
     "Cleat pockets hold one item each, and a blob count against an empty-belt photograph flags two items in one frame as mixed, so they go to Bin B for a hand count.")

h("4. How it works")
para("On patrol the boat follows its preset route between two barriers and returns to the dock when the bins are full or the battery is low. "
     "The upstream barrier is a 10 mm rope with a plastic float every 30 cm, strung at about 30 degrees to the flow between bank posts so mats and large debris slide to the bank. "
     "The downstream barrier is the same rope strung straight across, and a pond test checks sag and hold.")
para("A small INT8 object detector boxes each floating item and separates plants from plastic. The boat steers to the nearest plastic box and the scoop lifts it onto the belt, where the CNN names each item. "
     "Plastic goes to Bin A, and plants, mixed and low-confidence items go to Bin B. Training photographs are augmented with color shifts, blur, rotation, crops, random erasing and synthetic mud and droplets.")
para("Belt and thruster current is monitored. A jam triggers a reverse, up to three retries, then a stop and a return to the dock, and the drivers default to off. "
     "Counts become estimated kilograms and value, corrected by a weekly weigh-in, in an evidence sheet for the barangay and the DENR. "
     "On the bank, three float switches on a staff gauge set watch, alert and danger levels, and a buzzer and a small LED alert residents nearby while the level is logged.")

h("5. Power, sleep and safety")
fig("fig_wiring.png", 6.6, "Block schematic of the boat and the flood node (schematic).")
para("A 30 W panel yields about 54 Wh a day and a 20-minute patrol uses about 25 Wh. The 12 V 12 Ah battery stores about 144 Wh, with about 72 Wh usable, and we will measure these figures on the bench. "
     "Each motor has its own fuse and driver, the master switch cuts everything, a stop switch cuts motor power and a low-battery rule skips patrols.")
para("A camera board draws about 180 mA awake and 6 mA asleep at 5 V (datasheet values), about 22 Wh a day against 0.7 Wh, so sleep protects the 54 Wh solar budget. "
     "Docked boards sleep and wake on a timer for each patrol, then check the flood node's last status. On patrol the bow camera idles until it sees plastic and the belt camera and LED power up only when an item arrives. "
     "The flood node sleeps until a float switch trips and wakes every ten minutes to report.")
para("PVC floats, plastic bins and stainless hardware are materials local junk shops take, foam is kept to the minimum, and the lead-acid battery goes back to a dealer. "
     "The flood node's protected 18650 cell sits in an IP67 box with overcharge and over-discharge protection.")

h("6. Materials and cost (estimated online-shop prices, to be confirmed)")
rows = [("Item", "PHP")] + list(C.BUDGET)
table(rows, [5.2, 1.2], bold_rows=(len(rows) - 3, len(rows) - 2, len(rows) - 1))
para("Materials: PVC floats and deck, PVC or plastic stock for the scoop and belt frame, stainless fasteners (316 if the estero is brackish), IP65 boxes for the electronics, battery and bow camera and an IP67 box for the flood node cell. "
     "These are pilot-grade and should last months. At scale we would standardize to HDPE pipe floats and sturdier belts. Tools to borrow: soldering iron, multimeter, drill, saw, pliers, kitchen scale.")

h("7. What we expect (to be tested, not promised)")
para("Removal: at least 50% of the floating debris entering a 30 m patrol stretch against a baseline week, at least 20 kg lifted in 8 weeks and at least 1,000 classified items. "
     "Counting: weekly counts within 15% of a trained hand audit, and at least 85% precision and recall on photographs from unseen places. Flood alarm: sounds at all three marks in a bucket test.")
para("Limits: people must still empty the bins, clear snags and wipe the lenses, and we will log those hours. The boat covers one creek stretch and cannot lift heavy or tangled items.", size=10)

h("8. Build plan")
para("Nov 2: entry due. Nov 9 to 13: screening. If shortlisted, Nov 16 to Dec 5: finish Stage 1 on a bench (conveyor, camera, CNN, battery, panel), build a scaled boat mock-up and the flood node, and refine the 3D model (SAGIP_Boat_3D_Model.stl in Tinkercad) and the wiring schematic. "
     "Dec 7: 30-minute pitch with a bench demo. Creek pilot January to February 2027 if supported. Stage gates: model reaches 85% on held-out photos, float and tilt tests pass, the boat reaches and scoops at least 70% of 20 floating items in a pond, "
     "and the flood alarm sounds at all three marks.")
para("Sources: David et al. (2021), Proceedings of Machine Learning and Systems 3. Gholami et al. (2021), arXiv:2103.13630. Schreyers et al. (2021), Frontiers in Environmental Science 9, 686334. van Emmerik et al. (2022), Earth's Future 10, e2022EF002811. Full list in the written entry.", size=9)
d.save(OUT)
print("saved", OUT)
