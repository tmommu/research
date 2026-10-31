import re, sys, copy
sys.path.insert(0, '.')
import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_COLOR_INDEX, WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
import content as C

TPL = "templates/SFT_Entry_Template_official.docx"
OUT = sys.argv[1] if len(sys.argv) > 1 else "SAGIP_SFT2026_Proposal.docx"
d = docx.Document(TPL)
body = d.element.body
B = list(body)            # template body children, indexed as inspected
TOK = re.compile(r'(\*\*.+?\*\*|\[[^\]]+\])')
LINE = 360                # 1.5 spacing as the template reminders require

def clear_runs(p):
    for r in list(p._p):
        if r.tag in (qn('w:r'), qn('w:hyperlink')): p._p.remove(r)

def add_text(p, text, bold=None):
    for part in TOK.split(text):
        if not part: continue
        if part.startswith('**'):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith('['):
            r = p.add_run(part); r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            if bold: r.bold = True
        else:
            r = p.add_run(part)
            if bold: r.bold = True

def set_spacing(p, line=LINE, after=None, before=None):
    pPr = p._p.get_or_add_pPr()
    sp = pPr.find(qn('w:spacing'))
    if sp is None:
        sp = OxmlElement('w:spacing')
        # schema order: spacing precedes ind, jc, rPr
        anchor = None
        for tag in ('w:ind', 'w:contextualSpacing', 'w:jc', 'w:rPr'):
            anchor = pPr.find(qn(tag))
            if anchor is not None: break
        if anchor is not None: anchor.addprevious(sp)
        else: pPr.append(sp)
    sp.set(qn('w:line'), str(line)); sp.set(qn('w:lineRule'), 'auto')
    if after is not None: sp.set(qn('w:after'), str(after))
    if before is not None: sp.set(qn('w:before'), str(before))

def fill(p, label, text, label_bold=True):
    """Replace a template paragraph's runs, keeping its native pPr (numbering, indent)."""
    clear_runs(p)
    if label:
        r = p.add_run(label)
        if label_bold: r.bold = True
    if text: add_text(p, text)
    set_spacing(p)
    return p

def clone_after(ref, src, label, text, label_bold=True, strip_num=False, left=None):
    el = copy.deepcopy(src._p if isinstance(src, Paragraph) else src)
    ref_el = ref._p if isinstance(ref, Paragraph) else ref
    ref_el.addnext(el)
    p = Paragraph(el, d._body)
    if strip_num:
        pPr = el.find(qn('w:pPr'))
        n = pPr.find(qn('w:numPr'))
        if n is not None: pPr.remove(n)
        ind = pPr.find(qn('w:ind'))
        if ind is not None: pPr.remove(ind)
        if left is not None:
            ind = OxmlElement('w:ind'); ind.set(qn('w:left'), str(left)); ind.set(qn('w:firstLine'), '0')
            sp = pPr.find(qn('w:spacing'))
            (sp.addnext(ind) if sp is not None else pPr.insert(0, ind))
    fill(p, label, text, label_bold)
    return p

def P(i): return Paragraph(B[i], d._body)

def borders(tbl):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '8'); e.set(qn('w:space'), '0'); e.set(qn('w:color'), '000000')
        b.append(e)
    anchor = None
    for tag in ('w:tblInd', 'w:tblCellSpacing', 'w:jc', 'w:tblW', 'w:tblStyle'):
        anchor = tblPr.find(qn(tag))
        if anchor is not None: break
    if anchor is not None: anchor.addnext(b)
    else: tblPr.insert(0, b)

def make_table(after_el, rows, widths, bold_last=0):
    t = d.add_table(rows=len(rows), cols=len(widths))
    t.autofit = False
    borders(t)
    for ci, w in enumerate(widths): t.columns[ci].width = Inches(w)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = t.cell(ri, ci); c.width = Inches(widths[ci])
            p = c.paragraphs[0]
            p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
            if ri < len(rows) - 1: p.paragraph_format.keep_with_next = True
            isb = ri == 0 or (bold_last and ri >= len(rows) - bold_last)
            add_text(p, val, True if isb else None)
            if ci == len(widths) - 1 and re.fullmatch(r'[\d,]+', val or ''): p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    after_el.addnext(t._tbl)
    return t

def picture_after(ref, path, width, caption):
    el = OxmlElement('w:p'); ref_el = ref._p if isinstance(ref, Paragraph) else ref
    ref_el.addnext(el)
    p = Paragraph(el, d._body); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Inches(width))
    el2 = OxmlElement('w:p'); el.addnext(el2)
    c = Paragraph(el2, d._body); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(caption); r.italic = True
    return c

# ---------------- header table (template cells, labels kept) ----------------
T = d.tables[0]
REM = d.tables[1]._tbl
def add_after_label(cell, text):
    p = cell.paragraphs[0]
    r = p.add_run(' ' + text); r.bold = False
add_after_label(T.rows[0].cells[0], C.SCHOOL)
add_after_label(T.rows[0].cells[1], C.ADVISER)
c = T.rows[1].cells[0]
for i, m in enumerate(C.TEAM, 1):
    p = c.add_paragraph(); p.add_run(f"{i}. {m.split(',')[0]}")
c = T.rows[1].cells[1]
for i, m in enumerate(C.TEAM, 1):
    p = c.add_paragraph(); p.add_run(f"{i}. Grade 10")
c = T.rows[2].cells[0]
for p in c.paragraphs[2:]: p._element.getparent().remove(p._element)
c.paragraphs[1].runs[0].bold = True
c = T.rows[3].cells[0]
r = c.paragraphs[0].add_run(' ' + C.TITLE); r.bold = False
c = T.rows[4].cells[0]
p = c.add_paragraph(); add_text(p, C.DESC)

# ---------------- section 1 ----------------
p = fill(P(5), None, C.S1[0])
last = p
for t in C.S1[1:]:
    last = clone_after(last, P(5), None, t)
# ---------------- section 2 ----------------
fill(P(9), "Goal", ": " + C.GOAL)
obj = fill(P(10), "Objectives", " (SMART):")
last = obj
sub = P(27)                                   # template sub-bullet style (numId 3)
for lab, txt in C.OBJ:
    last = clone_after(last, sub, lab, txt)
# ---------------- section 3 ----------------
fill(P(14), None, C.S3)
# ---------------- section 4 ----------------
p = fill(P(18), "Select. ", C.S4_SELECT)
q = clone_after(p, P(18), "What exists. ", C.S4_EXISTS)
t = make_table(q._p, [("Existing work", "What it does", "What it leaves open")] + C.PRIOR, [2.1, 2.3, 2.1])
sp = OxmlElement('w:p'); t._tbl.addnext(sp)
g = clone_after(sp, P(18), "Gap we test. ", C.S4_GAP)
# ---------------- section 5 ----------------
p = fill(P(22), None, C.S5_STEM)
cap = picture_after(p, "fig_patrol.png", 6.0, "Figure 1. Patrol, detect, scoop and dock (top view, schematic).")
last = cap
for i, txt in enumerate(C.S5_PARAS):
    last = clone_after(last, P(18), None, txt)
    if i == C.S5_FIG2_AFTER:
        last = picture_after(last, "fig_boat_side.png", 6.4, "Figure 2. Side view of the SAGIP boat (schematic).")
last = clone_after(last, P(18), None, C.S5_PLAIN)
# ---------------- section 6 ----------------
fill(P(26), None, "The pilot unit costs about PHP " + C.BUDGET[-1][1] + " and is built in three stages, each bought only after the previous test passes.")
fill(P(27), "Resources: ", ""); P(27).paragraph_format.keep_with_next = True
t = make_table(P(27)._p, [("Item", "PHP")] + C.BUDGET, [5.4, 1.1], bold_last=3)
sp = OxmlElement('w:p'); t._tbl.addnext(sp)
last = clone_after(sp, P(26), "Tools. ", C.TOOLS)
for lab, txt in C.STAGES:
    last = clone_after(last, P(26), lab, txt)
last = clone_after(last, P(26), "Timeline. ", C.TIMELINE)
fill(P(28), "Stakeholders: ", "")
last = P(28)
for lab, txt in C.STAKE:
    last = clone_after(last, P(26), lab, txt)
fill(P(29), "Risks & Mitigation: ", "")
last = P(29)
for lab, txt in C.RISKS:
    last = clone_after(last, P(26), lab, txt)
# ---------------- section 7 ----------------
fill(P(33), None, C.S7_INTRO)
fill(P(35), "Measurable Outcomes:", "")
last = P(35)
for lab, txt in C.OUTCOMES:
    last = clone_after(last, P(26), lab, txt)
fill(P(36), "Long-Term Vision: ", C.VISION)
fill(P(37), "Key Takeaway: ", C.TAKEAWAY)
# fix labels in 36/37 template (label sits inside first run, no bold)
# ---------------- references ----------------
last = P(37)
h = clone_after(last, P(16), "References", None)
h._p.get_or_add_pPr()
n = h._p.pPr.find(qn('w:numPr'))
if n is not None: h._p.pPr.remove(n)
ind = h._p.pPr.find(qn('w:ind'))
if ind is not None: h._p.pPr.remove(ind)
last = h
for r in C.REFS:
    last = clone_after(last, P(26), None, r)
    set_spacing(last, line=276, after=60)
# ---------------- remove reminders table ----------------
REM.getparent().remove(REM)

d.save(OUT)

dd = docx.Document(OUT)
words = sum(len(p.text.split()) for p in dd.paragraphs)
for t in dd.tables:
    seen = set()
    for row in t.rows:
        for cell in row.cells:
            if id(cell._tc) in seen: continue
            seen.add(id(cell._tc))
            words += sum(len(p.text.split()) for p in cell.paragraphs)
print("words:", words)
